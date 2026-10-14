"""Import Taiwan dermatology board exam MCQs (101-115) into data/exam/questions.json.

Source: the Word files published by 臺灣皮膚科醫學會 (one question file and one
answer-key file per year). Statement markers and answer combinations are stored
as Wingdings symbols in those files, so this script decodes w:sym elements
instead of relying on paragraph text.

Usage:
  python _exam_import.py SOURCE_DIR [--extra-dir DIR] [--check]

--extra-dir points at a folder of .docx conversions for years whose files were
only provided as legacy .doc (named <year>_mcq_answers.docx /
<year>_image_answers.docx). For 115 the answer key ships as the misnamed
"115年度…【選擇題】.doc"; convert it with Word first. --check validates without
writing.

One-off maintenance tool, not part of the build: it needs python-docx
(pip install python-docx), unlike the stdlib-only pipeline scripts.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

import docx

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "data" / "exam"
YEARS = range(101, 116)
CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩"
# Simplified characters that slipped into the official Traditional Chinese files.
HANS_FIX = str.maketrans({"红": "紅", "脱": "脫"})
# Symbol-font private-use code points -> Unicode (Greek letters and common math signs).
SYMBOL_FONT = {
    **{0xF000 + ord(a): g for a, g in zip(
        "abgdezhqiklmnxoprstufcyw", "αβγδεζηθικλμνξοπρστυφχψω")},
    **{0xF000 + ord(a): g for a, g in zip(
        "ABGDEZHQIKLMNXOPRSTUFCYW", "ΑΒΓΔΕΖΗΘΙΚΛΜΝΞΟΠΡΣΤΥΦΧΨΩ")},
    0xF0B3: "≥", 0xF0A3: "≤", 0xF0B1: "±", 0xF0B4: "×", 0xF0AE: "→", 0xF0AD: "↑",
    0xF0AF: "↓", 0xF0AC: "←", 0xF0B0: "°", 0xF0BB: "≈", 0xF0B9: "≠", 0xF02D: "−",
    0xF0D7: "·", 0xF0A2: "′",
}


class ImportIssue(Exception):
    pass


def decode_sym(el, unknown: list[str]) -> str:
    font = (el.get(W + "font") or "").lower()
    raw = int(el.get(W + "char"), 16)
    # Symbol fonts store a single byte, either in the F0xx private-use range or
    # (in some converted files) as the cp1252 character for that byte.
    if raw >= 0xF000:
        byte = raw - 0xF000
    elif raw > 0xFF:
        try:
            byte = chr(raw).encode("cp1252")[0]
        except UnicodeEncodeError:
            byte = -1
    else:
        byte = raw
    if font == "wingdings" and 0x81 <= byte <= 0x8A:
        return CIRCLED[byte - 0x81]
    if font == "wingdings" and byte == 0xE0:
        return "→"
    if font == "wingdings 2" and 0x6A <= byte <= 0x73:
        return CIRCLED[byte - 0x6A]
    if font == "symbol" and 0xF000 + byte in SYMBOL_FONT:
        return SYMBOL_FONT[0xF000 + byte]
    unknown.append(f"{font}:{raw:X}")
    return "\N{REPLACEMENT CHARACTER}"


def para_text(p, unknown: list[str]) -> str:
    out = []
    for el in p._p.iter():
        if el.tag == W + "t":
            out.append(el.text or "")
        elif el.tag == W + "sym":
            out.append(decode_sym(el, unknown))
        elif el.tag == W + "tab":
            out.append(" ")
        elif el.tag in (W + "br", W + "cr"):
            out.append("\n")
    text = "".join(out).replace(" ", " ")
    # Dingbat circled digits (➀..➉) and math-italic letters (𝛼) -> plain forms.
    text = "".join(
        CIRCLED[ord(ch) - 0x2780] if 0x2780 <= ord(ch) <= 0x2789
        else unicodedata.normalize("NFKC", ch) if 0x1D400 <= ord(ch) <= 0x1D7FF
        else ch
        for ch in text
    )
    return re.sub(r"[ \t]+", " ", text).strip().translate(HANS_FIX)


def cell_paragraphs(cell, unknown: list[str]) -> list[str]:
    return [para_text(p, unknown) for p in cell.paragraphs]


CHOICE_LINE = re.compile(r"^\(?A\)?[.．、]")
CHOICE_SPLIT = re.compile(r"\(?([A-E])\)?[.．、]\s*")


def parse_choice_line(line: str) -> dict[str, str]:
    parts = CHOICE_SPLIT.split(line)
    # parts = ['', 'A', '①③ ', 'B', ...]
    choices = {}
    for i in range(1, len(parts) - 1, 2):
        choices[parts[i]] = parts[i + 1].replace("　", " ").strip()
    return choices


def parse_question(number: int, paragraphs: list[str]) -> dict:
    lines = [t for t in paragraphs if t]
    if not lines:
        raise ImportIssue(f"Q{number}: empty question cell")
    choice_idx = next((i for i in range(len(lines) - 1, -1, -1) if CHOICE_LINE.match(lines[i])), None)
    if choice_idx is None:
        raise ImportIssue(f"Q{number}: no A./B./... answer line")
    choice_line = " ".join(lines[choice_idx:])
    body = "\n".join(lines[:choice_idx])
    # Statements start at each circled numeral; some files put several on one
    # paragraph or append them to the stem after a line break.
    first = next((i for i, ch in enumerate(body) if ch in CIRCLED), None)
    stem_text = body if first is None else body[:first]
    stem_lines = [t.strip() for t in stem_text.split("\n") if t.strip()]
    statements = []
    if first is not None:
        for chunk in re.split(f"(?=[{CIRCLED}])", body[first:]):
            if chunk.strip():
                text = re.sub(r"\s*\n\s*", "\n", chunk[1:]).strip()
                statements.append({"id": chunk[0], "text": text})
        ids = [s["id"] for s in statements]
        if ids != list(CIRCLED[: len(ids)]):
            raise ImportIssue(f"Q{number}: statement markers out of order: {ids}")
    choices = parse_choice_line(choice_line)
    if not stem_lines:
        raise ImportIssue(f"Q{number}: missing stem")
    if not choices:
        raise ImportIssue(f"Q{number}: could not parse choices from {choice_line!r}")
    letters = "".join(choices)
    if letters != "ABCDE"[: len(letters)]:
        raise ImportIssue(f"Q{number}: choice letters out of order: {letters}")
    stmt_ids = {s["id"] for s in statements}
    for letter, combo in choices.items():
        if not combo:
            raise ImportIssue(f"Q{number}: choice {letter} is empty")
        missing = [ch for ch in combo if ch in CIRCLED and ch not in stmt_ids]
        if missing:
            raise ImportIssue(f"Q{number}: choice {letter} cites missing statement {missing}")
    # Single-statement choices (A=①, B=②, ...) read better as plain options.
    simple = statements and all(
        combo == statements[i]["id"] for i, combo in enumerate(choices.values())
    ) and len(choices) == len(statements)
    q = {"stem": "\n".join(stem_lines)}
    if simple:
        q["options"] = {letter: statements[i]["text"] for i, letter in enumerate(choices)}
    elif statements:
        q["statements"] = statements
        q["options"] = choices
    else:
        q["options"] = choices
    return q


def parse_question_file(path: Path, unknown: list[str]) -> dict[int, dict]:
    document = docx.Document(str(path))
    questions: dict[int, dict] = {}
    for table in document.tables:
        for row in table.rows:
            cells = row.cells
            first = para_text(cells[0].paragraphs[0], unknown) if cells[0].paragraphs else ""
            number_text = re.sub(r"\D", "", first)
            if not number_text:
                continue
            number = int(number_text)
            paragraphs = []
            seen = set()
            for i, cell in enumerate(cells):
                if id(cell._tc) in seen:
                    continue  # merged cells repeat
                seen.add(id(cell._tc))
                ps = cell_paragraphs(cell, unknown)
                if i == 0:
                    ps = ps[1:]
                paragraphs.extend(ps)
            if number in questions:
                raise ImportIssue(f"{path.name}: duplicate question {number}")
            questions[number] = parse_question(number, paragraphs)
    return questions


def parse_answer_file(path: Path, unknown: list[str]) -> dict[int, str]:
    document = docx.Document(str(path))
    answers: dict[int, str] = {}
    for table in document.tables:
        for row in table.rows:
            cells = []
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                cells.append([t for t in cell_paragraphs(cell, unknown) if t])
            for i in range(0, len(cells) - 1, 2):
                nums, ans = cells[i], cells[i + 1]
                if not nums or not all(re.fullmatch(r"\d{1,2}", n) for n in nums):
                    continue
                if len(nums) != len(ans):
                    raise ImportIssue(f"{path.name}: {nums} has {len(ans)} answers {ans}")
                for n, a in zip(nums, ans):
                    a = a.replace(" ", "").upper().replace("（", "(").replace("）", ")")
                    answers[int(n)] = a
    return answers


def normalize_answer(raw: str) -> list[str] | None:
    letters = re.findall(r"[A-E]", raw)
    return sorted(set(letters)) or None


def find(source: Path, extra: Path | None, year: int, kind: str) -> Path:
    names = {
        "questions": [f"{year}年度皮膚科專科醫師甄審筆試【選擇題】.docx"],
        "answers": [f"{year}年度皮膚科專科醫師甄審筆試【選擇題答案】.docx", f"{year}_mcq_answers.docx"],
    }[kind]
    for base in (source, extra):
        if base is None:
            continue
        for name in names:
            if (base / name).exists():
                return base / name
    raise ImportIssue(f"{year}: missing {kind} file ({names[0]})")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--extra-dir", type=Path)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()

    out = []
    problems: list[str] = []
    for year in YEARS:
        unknown: list[str] = []
        try:
            questions = parse_question_file(find(args.source, args.extra_dir, year, "questions"), unknown)
            answers = parse_answer_file(find(args.source, args.extra_dir, year, "answers"), unknown)
        except ImportIssue as exc:
            problems.append(f"{year}: {exc}")
            continue
        if sorted(questions) != list(range(1, 51)):
            problems.append(f"{year}: question numbers {sorted(questions)[:3]}..{len(questions)} != 1..50")
        if sorted(answers) != list(range(1, 51)):
            problems.append(f"{year}: answer numbers cover {len(answers)} of 50")
        if unknown:
            problems.append(f"{year}: unknown symbols {sorted(set(unknown))}")
        for number in sorted(questions):
            q = questions[number]
            raw = answers.get(number, "")
            key = normalize_answer(raw)
            if key is None or any(k not in q["options"] for k in key) or len(key) != 1:
                problems.append(f"{year}-{number:02d}: answer key {raw!r} needs review")
            if "送分" in raw or "給分" in raw:
                q["note"] = "官方公告此題送分"
            out.append({
                "id": f"{year}-{number:02d}",
                "year": year,
                "number": number,
                **q,
                "answer": key or [],
                "answerRaw": raw,
            })

    print(f"parsed {len(out)} questions from {len({q['year'] for q in out})} years")
    for p in problems:
        print("ISSUE", p)
    if not args.check:
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        path = OUT_DIR / "questions.json"
        path.write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
        print(f"wrote {path.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
